import io
import csv
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from datetime import datetime
from typing import Optional
from pydantic import BaseModel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from ..database import (
    progress_col, employees_col, onboarding_plans_col, validation_results_col,
    documents_col, chunks_col, requirement_matrix_col, audit_log_col,
)
from ..auth import require_roles

router = APIRouter(tags=["progress-reports"])


class ProgressUpdate(BaseModel):
    employee_id: str
    item_type: str  # module | checklist | task | quiz | assessment
    item_id: str
    status: Optional[str] = "Completed"
    score: Optional[float] = None


@router.post("/progress/update")
async def update_progress(payload: ProgressUpdate,
                           user=Depends(require_roles("admin", "training_manager", "employee"))):
    if user["role"] == "employee" and user.get("employee_id") != payload.employee_id:
        raise HTTPException(status_code=403, detail="You can only update your own progress")
    field_map = {
        "module": "module_completion", "checklist": "checklist_completion",
        "task": "task_completion", "quiz": "quiz_scores", "assessment": "assessment_scores",
    }
    field = field_map.get(payload.item_type)
    if not field:
        raise HTTPException(status_code=400, detail="item_type must be one of module/checklist/task/quiz/assessment")

    value = payload.score if payload.score is not None else payload.status
    await progress_col.update_one(
        {"employee_id": payload.employee_id},
        {"$set": {f"{field}.{payload.item_id}": value, "updated_at": datetime.utcnow().isoformat()}},
        upsert=True,
    )
    return {"message": "Progress updated"}


@router.get("/progress/{employee_id}")
async def get_progress(employee_id: str, user=Depends(require_roles(
        "admin", "training_manager", "reviewer", "manager", "employee"))):
    if user["role"] == "employee" and user.get("employee_id") != employee_id:
        raise HTTPException(status_code=403, detail="You can only view your own progress")
    progress = await progress_col.find_one({"employee_id": employee_id}, {"_id": 0})
    if not progress:
        raise HTTPException(status_code=404, detail="No progress record found")
    return progress


@router.get("/dashboard/admin")
async def admin_dashboard(user=Depends(require_roles("admin", "training_manager"))):
    total_employees = await employees_col.count_documents({})
    total_docs = await documents_col.count_documents({})
    total_plans = await onboarding_plans_col.count_documents({})
    validations = await validation_results_col.find({}, {"_id": 0}).to_list(length=2000)
    status_breakdown = {}
    for v in validations:
        status_breakdown[v["final_status"]] = status_breakdown.get(v["final_status"], 0) + 1
    avg_coverage = round(sum(v["coverage"]["coverage_score"] for v in validations) / len(validations), 2) if validations else 0
    avg_traceability = round(sum(v["traceability"]["traceability_score"] for v in validations) / len(validations), 2) if validations else 0
    return {
        "total_employees": total_employees,
        "total_documents": total_docs,
        "total_plans_generated": total_plans,
        "status_breakdown": status_breakdown,
        "average_coverage_score": avg_coverage,
        "average_traceability_score": avg_traceability,
    }


class PolicyUpdateImpactRequest(BaseModel):
    document_id: str


@router.post("/policy-update/impact-analysis")
async def policy_update_impact(payload: PolicyUpdateImpactRequest,
                                user=Depends(require_roles("admin", "training_manager"))):
    """Step 57/58/59 - Policy Update Detection + Impact Analysis + Selective Regeneration target list."""
    doc = await documents_col.find_one({"document_id": payload.document_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    affected_requirements = await requirement_matrix_col.find(
        {"source_document_id": payload.document_id}, {"_id": 0}
    ).to_list(length=2000)
    affected_req_ids = {r["requirement_id"] for r in affected_requirements}
    affected_roles = {r["role_name"] for r in affected_requirements}

    affected_employees = await employees_col.find(
        {"role_name": {"$in": list(affected_roles)}}, {"_id": 0}
    ).to_list(length=2000)

    all_plans = await onboarding_plans_col.find(
        {"employee_id": {"$in": [e["employee_id"] for e in affected_employees]}}, {"_id": 0}
    ).to_list(length=2000)

    affected_plan_ids = []
    for plan_rec in all_plans:
        plan = plan_rec["plan"]
        items = []
        for key in ("modules", "checklist", "tasks", "quizzes", "assessments"):
            items.extend(plan.get(key, []) or [])
        if any(i.get("requirement_id") in affected_req_ids for i in items):
            affected_plan_ids.append(plan_rec["plan_id"])

    await audit_log_col.insert_one({
        "action": "policy_update_impact_analysis", "document_id": payload.document_id,
        "affected_requirement_count": len(affected_req_ids),
        "affected_employee_count": len(affected_employees),
        "affected_plan_count": len(affected_plan_ids),
        "user": user["username"], "timestamp": datetime.utcnow().isoformat(),
    })

    return {
        "document": doc,
        "affected_requirements": list(affected_req_ids),
        "affected_roles": list(affected_roles),
        "affected_employees": [e["employee_id"] for e in affected_employees],
        "plans_requiring_regeneration": affected_plan_ids,
        "recommendation": "Call /onboarding/generate again for each affected employee_id "
                           "(selective regeneration - only these plans, not the entire system).",
    }


@router.get("/reports/coverage.csv")
async def export_coverage_csv(user=Depends(require_roles("admin", "training_manager", "reviewer"))):
    validations = await validation_results_col.find({}, {"_id": 0}).to_list(length=5000)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["employee_id", "final_status", "coverage_score", "traceability_score",
                      "missing_requirement_count", "unsupported_requirement_count", "contradiction_count"])
    for v in validations:
        writer.writerow([
            v.get("employee_id"), v.get("final_status"),
            v["coverage"]["coverage_score"], v["traceability"]["traceability_score"],
            v["consistency_count"]["missing_requirement_count"],
            v["consistency_count"]["unsupported_requirement_count"],
            v["consistency_count"]["contradiction_count"],
        ])
    buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                              headers={"Content-Disposition": "attachment; filename=coverage_report.csv"})


@router.get("/reports/coverage.xlsx")
async def export_coverage_xlsx(user=Depends(require_roles("admin", "training_manager", "reviewer"))):
    """Step 63 - Export in Excel-compatible format."""
    validations = await validation_results_col.find({}, {"_id": 0}).to_list(length=5000)

    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage Report"

    headers = ["Employee ID", "Final Status", "Coverage Score (%)", "Traceability Score (%)",
               "Missing Requirements", "Unsupported Requirements", "Contradictions", "Validated At"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="C9A227", end_color="C9A227", fill_type="solid")

    for v in validations:
        ws.append([
            v.get("employee_id"), v.get("final_status"),
            v["coverage"]["coverage_score"], v["traceability"]["traceability_score"],
            v["consistency_count"]["missing_requirement_count"],
            v["consistency_count"]["unsupported_requirement_count"],
            v["consistency_count"]["contradiction_count"],
            v.get("validated_at"),
        ])

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=coverage_report.xlsx"},
    )


@router.get("/reports/coverage.pdf")
async def export_coverage_pdf(user=Depends(require_roles("admin", "training_manager", "reviewer"))):
    """Step 63 - Export in PDF format."""
    validations = await validation_results_col.find({}, {"_id": 0}).to_list(length=5000)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    elements = [Paragraph("SkillSprint AI — Coverage & Validation Report", styles["Title"]),
                Paragraph(f"Generated: {datetime.utcnow().isoformat()}", styles["Normal"]),
                Spacer(1, 16)]

    data = [["Employee ID", "Final Status", "Coverage %", "Traceability %", "Missing", "Unsupported", "Contradictions"]]
    for v in validations:
        data.append([
            v.get("employee_id"), v.get("final_status"),
            v["coverage"]["coverage_score"], v["traceability"]["traceability_score"],
            v["consistency_count"]["missing_requirement_count"],
            v["consistency_count"]["unsupported_requirement_count"],
            v["consistency_count"]["contradiction_count"],
        ])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C9A227")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (2, 1), (-1, -1), "CENTER"),
    ]))
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
                              headers={"Content-Disposition": "attachment; filename=coverage_report.pdf"})


@router.get("/reports/requirement-matrix.xlsx")
async def export_requirement_matrix_xlsx(user=Depends(require_roles("admin", "training_manager", "reviewer"))):
    """Full Role Requirement Matrix as an Excel-compatible export."""
    reqs = await requirement_matrix_col.find({}, {"_id": 0}).to_list(length=5000)

    wb = Workbook()
    ws = wb.active
    ws.title = "Role Requirement Matrix"

    if reqs:
        headers = list(reqs[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="C9A227", end_color="C9A227", fill_type="solid")
        for r in reqs:
            ws.append([str(r.get(h, "")) for h in headers])
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=role_requirement_matrix.xlsx"},
    )
